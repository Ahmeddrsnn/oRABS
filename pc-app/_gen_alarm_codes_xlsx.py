"""Alarm kodu listesini Excel'e döker (tek kaynak: core/alarm_codes.py).

Komut: python _gen_alarm_codes_xlsx.py
Çıktı: Alarm_Kodlari.xlsx (bu dizinde)
"""
from __future__ import annotations

import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, 'src')
if SRC not in sys.path:
    sys.path.insert(0, SRC)

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.alarm_codes import all_rows

GROUP_NAMES = {
    'KP': 'Kapı',
    'MS': 'Makine Durdurma',
    'IB': 'Işık Bariyeri',
    'AS': 'Acil Stop',
    'PT': 'Port',
    'LF': 'LAF Ünitesi',
    'AI': 'Analog Giriş',
    'BT': 'Batch',
}
SEV_TR = {'crit': 'KRİTİK', 'warn': 'UYARI', 'info': 'BİLGİ'}
SEV_FILL = {
    'crit': PatternFill('solid', fgColor='FEE2E2'),
    'warn': PatternFill('solid', fgColor='FEF3C7'),
    'info': PatternFill('solid', fgColor='DBEAFE'),
}

wb = Workbook()
ws = wb.active
ws.title = 'Alarm Kodları'

headers = ['Kod', 'Grup', 'Açıklama', 'Önem', 'Sistem Olayı (action_type)']
ws.append(headers)
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.font = Font(bold=True, color='FFFFFF', size=12)
    cell.fill = PatternFill('solid', fgColor='1F2937')
    cell.alignment = Alignment(horizontal='center', vertical='center')

for code, label, sev, action in all_rows():
    grp = GROUP_NAMES.get(code.split('-')[0], '')
    ws.append([code, grp, label, SEV_TR.get(sev, sev), action])
    r = ws.max_row
    ws.cell(row=r, column=1).font = Font(bold=True)
    ws.cell(row=r, column=4).fill = SEV_FILL.get(sev, PatternFill())
    ws.cell(row=r, column=4).alignment = Alignment(horizontal='center')

widths = [10, 18, 52, 12, 38]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:E{ws.max_row}'

out = os.path.join(HERE, 'Alarm_Kodlari.xlsx')
wb.save(out)
print(f'YAZILDI: {out}  ({ws.max_row - 1} kod)')
